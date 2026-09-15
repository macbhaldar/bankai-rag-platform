"""File loaders for the ingestion pipeline (markdown, txt, PDF, DOCX, CSV, JSON)."""
from __future__ import annotations
 
import csv
import io
import json
import re
from dataclasses import dataclass, field
from pathlib import Path
 
SUPPORTED_EXTS = {".md", ".txt", ".pdf", ".docx", ".csv", ".json"}
 
 
@dataclass
class LoadedDocument:
    path: Path
    filename: str
    stem: str
    text: str
    n_pages: int = 1
    warnings: list[str] = field(default_factory=list)
 
 
def discover_files(root: Path) -> list[Path]:
    if root.is_file():
        return [root]
    return sorted(
        p for p in root.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED_EXTS
    )
 
 
def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)  # de-hyphenate line breaks
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
 
 
def _load_csv(path: Path) -> str:
    with path.open("r", encoding="utf-8", errors="replace", newline="") as fh:
        rows = list(csv.reader(fh))
    if not rows:
        return ""
    rows = rows[:500]
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerows(rows)
    if len(rows) == 500:
        buf.write("\n... (truncated at 500 rows)")
    return buf.getvalue()
 
 
def load_file(path: Path) -> LoadedDocument:
    ext = path.suffix.lower()
    warnings: list[str] = []
    if ext in {".md", ".txt"}:
        text = path.read_text(encoding="utf-8", errors="replace")
        pages = 1
    elif ext == ".pdf":
        from pypdf import PdfReader
 
        reader = PdfReader(str(path))
        page_texts = [(page.extract_text() or "") for page in reader.pages]
        text = "\n\n".join(page_texts)
        pages = len(reader.pages)
    elif ext == ".docx":
        import docx  # python-docx
 
        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(parts)
        pages = 1
    elif ext == ".csv":
        text = _load_csv(path)
        pages = 1
    elif ext == ".json":
        text = json.dumps(json.loads(path.read_text(encoding="utf-8")), indent=2)
        pages = 1
    else:
        raise ValueError(f"unsupported file type: {ext}")
    if not text.strip():
        warnings.append("empty document")
    return LoadedDocument(
        path=path,
        filename=path.name,
        stem=path.stem,
        text=clean_text(text),
        n_pages=pages,
        warnings=warnings,
    )